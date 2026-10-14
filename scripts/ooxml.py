"""python-docx helpers for rebuilding a PDF's layout and styling in Word.

python-docx exposes no API for run/paragraph shading, character spacing,
paragraph borders, page borders or floating (page-anchored) pictures.  All of
those are raw OOXML.  Word is strict about the order of children inside
``pPr``/``rPr``/``tblPr``/``tcPr``/``sectPr``, so every insertion goes through
``insert_ordered()`` instead of plain ``append()``.

Typical use: write a per-document build script that imports this module, then
call the helpers below while walking the source PDF's content.
"""

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

CHILD_ORDER = {
    "pPr": [
        "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
        "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
        "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
        "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
        "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
        "suppressOverlap", "jc", "textDirection", "textAlignment",
        "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
        "pPrChange",
    ],
    "rPr": [
        "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
        "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
        "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
        "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
        "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
        "eastAsianLayout", "specVanish", "oMath",
    ],
    "tblPr": [
        "tblStyle", "tblpPr", "tblOverlap", "bidiVisual",
        "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
        "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
        "tblCellMar", "tblLook", "tblCaption", "tblDescription",
    ],
    "tcPr": [
        "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
        "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    ],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore",
             "wAfter", "cantSplit", "trHeight", "tblHeader", "tblCellSpacing",
             "jc", "hidden"],
    "sectPr": [
        "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
        "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
        "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter",
        "docGrid", "printerSettings",
    ],
}


def insert_ordered(parent, child, kind=None):
    """Insert ``child`` into ``parent`` at its schema-mandated position."""
    kind = kind or parent.tag.split("}")[1]
    order = CHILD_ORDER.get(kind)
    if order:
        tag = child.tag.split("}")[1]
        index = order.index(tag)
        for existing in parent:
            etag = existing.tag.split("}")[1]
            if etag in order and order.index(etag) > index:
                existing.addprevious(child)
                return child
    parent.append(child)
    return child


# --------------------------------------------------------------------- setup
def new_document(page_cm=(21.0, 29.7), margins_cm=(1.1, 0.7, 1.45, 1.75),
                 font="Microsoft YaHei", east_asian="微软雅黑", latin=None,
                 size=10.5, color="414141", line_pt=18.0):
    """Create an A4-style document with a Chinese-document Normal style.

    ``margins_cm`` is ``(top, right, bottom, left)``.  ``latin`` overrides the
    font used for ASCII/Latin text (defaults to ``font``).  ``east_asian`` is
    required: Word otherwise falls back to the theme's East Asian font and the
    Chinese text comes out in the wrong typeface.
    """
    from docx import Document

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(page_cm[0])
    section.page_height = Cm(page_cm[1])
    top, right, bottom, left = margins_cm
    section.top_margin = Cm(top)
    section.right_margin = Cm(right)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)

    normal = doc.styles["Normal"]
    normal.font.name = latin or font
    normal.font.size = Pt(size)
    normal.font.color.rgb = RGBColor.from_string(color)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr, value in (("w:ascii", latin or font), ("w:hAnsi", latin or font),
                        ("w:eastAsia", east_asian), ("w:cs", latin or font)):
        rFonts.set(qn(attr), value)

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if line_pt:
        pf.line_spacing = Pt(line_pt)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    return doc


# -------------------------------------------------------------------- runs
def style_run(run, size=10.5, bold=False, color="414141", font=None,
              east_asian="微软雅黑", spacing_pt=None, shading=None,
              underline=False, all_caps=False):
    """Apply size/colour/font plus optional character spacing and highlight.

    ``spacing_pt`` adds tracking after every character (use it to emulate a
    "distributed" label field, e.g. spread 姓名 over three full-width cells).
    ``shading`` paints a background behind the glyphs only, which is the
    closest text-based equivalent of a coloured heading ribbon.
    """
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if underline:
        run.font.underline = True
    if all_caps:
        run.font.all_caps = True
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = insert_ordered(rPr, OxmlElement("w:rFonts"))
    for attr, value in (("w:ascii", font), ("w:hAnsi", font),
                        ("w:eastAsia", east_asian), ("w:cs", font)):
        if value:
            rFonts.set(qn(attr), value)
    if spacing_pt:
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(int(round(spacing_pt * 20))))
        insert_ordered(rPr, el)
    if shading:
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:color"), "auto")
        el.set(qn("w:fill"), shading)
        insert_ordered(rPr, el)
    return run


def add_run(paragraph, text, **kwargs):
    return style_run(paragraph.add_run(text), **kwargs)


def symbol_run(paragraph, char, font="Wingdings", size=10.5, color="414141"):
    """Insert a symbol-font character (Word stores bullets as U+F0xx codes)."""
    run = style_run(paragraph.add_run(char), size=size, color=color, font=font,
                    east_asian=font)
    rFonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    rFonts.set(qn("w:hint"), "default")
    return run


def add_hyperlink(paragraph, url, text, size=10.5, color="0000FF", font=None,
                  east_asian="微软雅黑", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = paragraph.add_run(text)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)
    style_run(run, size=size, color=color, font=font, east_asian=east_asian,
              underline=underline)
    return run


# -------------------------------------------------------------- paragraphs
def set_para(paragraph, line=None, before=0, after=0, left=None, right=None,
             hanging=None, first_line=None, align=None, tabs=None,
             keep_with_next=False, keep_lines=False):
    """Spacing/indent/tab control.

    ``line`` is an exact line height in points (the norm for Chinese business
    documents, e.g. 10.5pt text on an 18pt rhythm).  ``left``/``right`` are cm
    and may be negative to pull a heading band outside the text margin.
    ``hanging`` is a cm hanging indent for bullets.
    """
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if left is not None:
        pf.left_indent = Cm(left)
    if right is not None:
        pf.right_indent = Cm(right)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    elif first_line is not None:
        pf.first_line_indent = Cm(first_line)
    if align is not None:
        pf.alignment = align
    pPr = paragraph._p.get_or_add_pPr()
    if keep_with_next:
        insert_ordered(pPr, OxmlElement("w:keepNext"))
    if keep_lines:
        insert_ordered(pPr, OxmlElement("w:keepLines"))
    if tabs:
        tabs_el = OxmlElement("w:tabs")
        for pos in tabs:
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "left")
            tab.set(qn("w:pos"), str(int(round(pos * 567))))
            tabs_el.append(tab)
        insert_ordered(pPr, tabs_el)
    return paragraph


def para_border(paragraph, side="bottom", color="4E7282", sz=6, space=0,
                val="single"):
    """Add one paragraph border.  ``sz`` is in eighths of a point (6 = 0.75pt).

    A full-width bottom border on a heading paragraph reproduces the horizontal
    rule that usually sits under a section ribbon in the source PDF.
    """
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = insert_ordered(pPr, OxmlElement("w:pBdr"))
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)
    return paragraph


def para_shading(paragraph, fill):
    """Fill the whole line box (contrast with run-level ``shading``)."""
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    insert_ordered(pPr, el)
    return paragraph


def page_border(section, sides=("left",), color="4E7282", sz=6, space=9,
                offset_from="text"):
    """Draw page borders, e.g. the thin vertical accent bar along the margin."""
    sectPr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), offset_from)
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        borders.append(el)
    insert_ordered(sectPr, borders)
    return section


# ------------------------------------------------------------------ tables
def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    insert_ordered(tcPr, el)
    return cell


def table_borders(table, sides=("top", "left", "bottom", "right", "insideH",
                                "insideV"), val="none", color="auto", sz=0):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    insert_ordered(tblPr, borders)
    return table


def cell_margins(table, top=0, left=0, bottom=0, right=0):
    mar = OxmlElement("w:tblCellMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom),
                        ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(int(value)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    insert_ordered(table._tbl.tblPr, mar)
    return table


def fixed_table(table, widths_cm, total_cm=None, indent_cm=None):
    """Pin column widths (Word ignores them otherwise) and drop the borders."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    insert_ordered(tblPr, layout)
    total = total_cm if total_cm is not None else sum(widths_cm)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(round(total * 567))))
    w.set(qn("w:type"), "dxa")
    insert_ordered(tblPr, w)
    if indent_cm is not None:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(int(round(indent_cm * 567))))
        ind.set(qn("w:type"), "dxa")
        insert_ordered(tblPr, ind)
    table_borders(table)
    cell_margins(table)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    return table


def row_height(row, pt, rule="exact"):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(round(pt * 20))))
    el.set(qn("w:hRule"), rule)
    insert_ordered(trPr, el)
    return row


def cell_valign(cell, value="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:vAlign")
    el.set(qn("w:val"), value)
    insert_ordered(tcPr, el)
    return cell


# ------------------------------------------------------------------ images
def add_floating_picture(paragraph, image_path, x_pt, y_pt, w_pt, h_pt,
                         behind_doc=False):
    """Place a picture at an absolute position on the page.

    ``x_pt``/``y_pt`` are measured from the top-left corner of the physical
    page, so coordinates read straight out of the source PDF can be reused.
    IMPORTANT: ``paragraph`` must be a body-level paragraph.  When the anchor
    lives inside a table cell Word adds the cell's origin to the offset, which
    shifts the picture by the page margins.
    """
    run = paragraph.add_run()
    run.add_picture(image_path, width=Pt(w_pt), height=Pt(h_pt))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    for key, value in (
        ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
        ("simplePos", "0"), ("relativeHeight", "2"),
        ("behindDoc", "1" if behind_doc else "0"), ("locked", "0"),
        ("layoutInCell", "1"), ("allowOverlap", "1"),
    ):
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for tag, offset in (("wp:positionH", x_pt), ("wp:positionV", y_pt)):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset")
        off.text = str(int(Emu(Pt(offset))))
        pos.append(off)
        anchor.append(pos)
    anchor.append(inline.find(qn("wp:extent")))
    effect = inline.find(qn("wp:effectExtent"))
    if effect is None:
        effect = OxmlElement("wp:effectExtent")
        for side in ("l", "t", "r", "b"):
            effect.set(side, "0")
    anchor.append(effect)
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(inline.find(qn(tag)))
    drawing.replace(inline, anchor)
    return run


def spacer(doc, size=1.0):
    """Tiny empty paragraph.  Required between two adjacent tables: Word merges
    consecutive tables that are not separated by a paragraph."""
    p = doc.add_paragraph()
    set_para(p, line=max(size, 1.0))
    add_run(p, "", size=size)
    return p
